"""Excel template and import helpers using openpyxl."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

THUMB_W = 96
THUMB_H = 72


def _workbook_response(wb: Workbook, filename: str) -> HttpResponse:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def build_template_response(
    filename: str,
    headers: list[str],
    sample_rows: list[list[Any]] | None = None,
    sheet_title: str = "Data",
) -> HttpResponse:
    """Create an .xlsx download with header row and optional sample rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill = PatternFill("solid", fgColor="1A3A5C")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 2)

    if sample_rows:
        for r_idx, row in enumerate(sample_rows, start=2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

    return _workbook_response(wb, filename)


def _make_excel_thumbnail(path: str | Path) -> BytesIO | None:
    """Small JPEG thumbnail so Excel embeds stay row-sized."""
    try:
        from PIL import Image, ImageOps
    except ImportError:
        return None

    image_file = Path(path)
    if not image_file.is_file():
        return None
    try:
        with Image.open(image_file) as image:
            image = ImageOps.exif_transpose(image)
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            image.thumbnail((THUMB_W, THUMB_H), Image.Resampling.LANCZOS)
            buf = BytesIO()
            image.save(buf, format="JPEG", quality=75, optimize=True)
            buf.seek(0)
            buf.name = "thumb.jpg"
            return buf
    except Exception:  # noqa: BLE001
        return None


def _add_row_image(ws, xl_img: XLImage, col_1based: int, row_1based: int) -> None:
    """Pin image to a single cell so it stays with that data row."""
    marker = AnchorMarker(col=col_1based - 1, colOff=0, row=row_1based - 1, rowOff=0)
    size = XDRPositiveSize2D(pixels_to_EMU(THUMB_W), pixels_to_EMU(THUMB_H))
    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(xl_img)


def build_data_export_response(
    filename: str,
    headers: list[str],
    rows: Iterable[list[Any]],
    *,
    sheet_title: str = "Data",
    image_paths: list[str | None] | None = None,
    image_column_header: str = "Photo",
) -> HttpResponse:
    """
    Export data rows to .xlsx.

    When image_paths is provided (same length as rows), adds a Photo column and
    embeds a fixed-size thumbnail pinned to that row.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill = PatternFill("solid", fgColor="1A3A5C")
    header_font = Font(color="FFFFFF", bold=True)
    export_headers = list(headers)
    if image_paths is not None:
        export_headers.append(image_column_header)

    for col, header in enumerate(export_headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        width = 14 if header == image_column_header else max(14, len(header) + 2)
        ws.column_dimensions[get_column_letter(col)].width = width

    photo_col = len(export_headers) if image_paths is not None else None
    row_list = list(rows)
    for r_idx, row in enumerate(row_list, start=2):
        for c_idx, value in enumerate(row, start=1):
            if value is None:
                cell_value = ""
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell_value = value
            else:
                cell_value = str(value)
            ws.cell(row=r_idx, column=c_idx, value=cell_value)

        if photo_col is None or image_paths is None:
            continue

        path = image_paths[r_idx - 2] if r_idx - 2 < len(image_paths) else None
        if not path:
            ws.cell(row=r_idx, column=photo_col, value="")
            continue

        image_file = Path(path)
        label = image_file.name if image_file.name else "Photo"
        thumb = _make_excel_thumbnail(image_file)
        if thumb is None:
            ws.cell(row=r_idx, column=photo_col, value=label)
            continue

        try:
            xl_img = XLImage(thumb)
            xl_img.width = THUMB_W
            xl_img.height = THUMB_H
            _add_row_image(ws, xl_img, photo_col, r_idx)
            ws.row_dimensions[r_idx].height = 58
            # Keep filename in cell for clarity / if image fails to render
            ws.cell(row=r_idx, column=photo_col, value=label)
        except Exception:  # noqa: BLE001
            ws.cell(row=r_idx, column=photo_col, value=label)

    return _workbook_response(wb, filename)


def read_sheet_rows(uploaded_file, expected_headers: list[str]) -> list[dict[str, Any]]:
    """
    Read first worksheet rows keyed by header name.

    Raises ValueError if required headers are missing.
    Skips completely empty rows.
    """
    wb = load_workbook(uploaded_file, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValueError("Excel file is empty.")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    missing = [h for h in expected_headers if h not in headers]
    if missing:
        raise ValueError(
            "Missing columns in Excel: "
            + ", ".join(missing)
            + ". Download the template and keep the header row."
        )

    index = {name: headers.index(name) for name in expected_headers}
    result: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        item = {}
        for name, col in index.items():
            value = row[col] if col < len(row) else None
            if isinstance(value, str):
                value = value.strip()
            elif isinstance(value, float) and value.is_integer():
                # Keep "80" as 80 (not 80.0) so free-text columns stay clean.
                value = int(value)
            item[name] = value
        result.append(item)
    return result


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
